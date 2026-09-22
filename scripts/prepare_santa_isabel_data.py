import openpyxl
import json
import uuid
import re
from datetime import datetime
from collections import defaultdict

NAMESPACE = uuid.NAMESPACE_DNS

def make_uuid(key: str) -> str:
    return str(uuid.uuid5(NAMESPACE, key))

def parse_date(d_str) -> str:
    if not d_str:
        return '2026-01-01'
    if isinstance(d_str, datetime):
        return d_str.strftime('%Y-%m-%d')
    s = str(d_str).strip().replace('\xa0', ' ')
    # Match day/month/year
    m = re.match(r'^(\d{1,2})/(\d{1,2})/(\d{4})', s)
    if m:
        day, month, year = int(m.group(1)), int(m.group(2)), int(m.group(3))
        return f'{year:04d}-{month:02d}-{day:02d}'
    # Match year-month-day
    m2 = re.match(r'^(\d{4})-(\d{1,2})-(\d{1,2})', s)
    if m2:
        year, month, day = int(m2.group(1)), int(m2.group(2)), int(m2.group(3))
        return f'{year:04d}-{month:02d}-{day:02d}'
    return '2026-01-01'

def parse_amount(a_str) -> float:
    if a_str is None:
        return 0.0
    if isinstance(a_str, (int, float)):
        return float(a_str)
    s = str(a_str).strip().replace(' ', '').replace('$', '')
    if ',' in s and '.' in s:
        s = s.replace('.', '').replace(',', '.')
    elif ',' in s:
        s = s.replace(',', '.')
    try:
        return round(float(s), 2)
    except:
        return 0.0

def clean_phone(val) -> str:
    if not val:
        return ''
    s = str(val).strip()
    s = re.sub(r'^[/\s]+|[/\s]+$', '', s)
    if '/' in s:
        parts = [p.strip() for p in s.split('/') if p.strip()]
        s = parts[0] if parts else ''
    s = re.sub(r'\s+', ' ', s).strip()
    return s[:20]

def clean_name(val) -> str:
    if not val:
        return 'Cliente General'
    s = re.sub(r'\s+', ' ', str(val)).strip()
    return s.title()

def clean_doc(val) -> str:
    if not val:
        return ''
    return re.sub(r'[^0-9A-Za-z-]', '', str(val)).strip()[:20]

def main():
    print("Loading Base de datos smart.xlsx...")
    wb_smart = openpyxl.load_workbook('Base de datos smart.xlsx', data_only=True)
    
    print("Loading Base de datos bitacora.xlsx...")
    wb_bit = openpyxl.load_workbook('Base de datos bitacora.xlsx', data_only=True)

    empresa_id = '00000000-0000-0000-0000-000000000001' # Inversiones Ingenieria GC S.A.S
    proyecto_id = '00000000-0000-0000-0000-000000000101' # Santa Isabel
    today = '2026-09-22'

    # Build bitacora client lookup for enrichment
    bit_clients = {}
    sheet_bit_cli = wb_bit['Cliente']
    rows_bit_cli = list(sheet_bit_cli.iter_rows(values_only=True))
    if rows_bit_cli:
        hdr_bc = rows_bit_cli[0]
        id_idx = hdr_bc.index('Identificacion')
        nom_idx = hdr_bc.index('ClienteNombre')
        cel_idx = hdr_bc.index('Celular')
        dir_idx = hdr_bc.index('Direccion')
        ciu_idx = hdr_bc.index('Ciudad')
        cor_idx = hdr_bc.index('Correo')
        for r in rows_bit_cli[1:]:
            doc = clean_doc(r[id_idx])
            if doc:
                bit_clients[doc] = {
                    'nombre': clean_name(r[nom_idx]),
                    'celular': clean_phone(r[cel_idx]),
                    'direccion': str(r[dir_idx]).strip() if r[dir_idx] else '',
                    'ciudad': str(r[ciu_idx]).strip() if r[ciu_idx] else '',
                    'correo': str(r[cor_idx]).strip() if r[cor_idx] else ''
                }

    # 1. Process Module -> Inmuebles
    print("Processing Module...")
    sheet_mod = wb_smart['Module']
    rows_mod = list(sheet_mod.iter_rows(values_only=True))
    hdr_mod = rows_mod[0]
    inm_idx = hdr_mod.index('Inmueble')
    val_tot_idx = hdr_mod.index('ValorTotal')
    val_inm_idx = hdr_mod.index('ValorInmueble')
    status_idx = hdr_mod.index('Status')

    # Also read DetalleCliente to get prices and status
    sheet_dc = wb_smart['DetalleCliente']
    rows_dc = list(sheet_dc.iter_rows(values_only=True))
    hdr_dc = rows_dc[0]
    dc_by_lot = {}
    for r in rows_dc[1:]:
        lot = str(r[hdr_dc.index('Modulo')]).strip()
        dc_by_lot[lot] = r

    inmuebles = []
    lot_to_inmueble_id = {}
    
    # Specific known prices for lots that were 0 in Module but priced in Bitacora
    fallback_lot_prices = {
        'LOTE 01': 44000000.0,
        'LOTE 155': 55000000.0,
    }

    for r in rows_mod[1:]:
        raw_lot = r[inm_idx]
        if not raw_lot:
            continue
        lot_name = str(raw_lot).strip()
        inm_id = make_uuid(f'inmueble_santa_isabel_{lot_name}')
        lot_to_inmueble_id[lot_name] = inm_id

        dc_row = dc_by_lot.get(lot_name)
        status_raw = str(r[status_idx]).strip() if r[status_idx] else 'Disponible'
        
        # Price determination
        price = 0.0
        if dc_row:
            price = parse_amount(dc_row[hdr_dc.index('ValorTotal')])
            if price <= 0:
                price = parse_amount(dc_row[hdr_dc.index('PrecioLista')])
        if price <= 0:
            price = parse_amount(r[val_tot_idx])
        if price <= 0:
            price = parse_amount(r[val_inm_idx])
        if price <= 0 and lot_name in fallback_lot_prices:
            price = fallback_lot_prices[lot_name]

        # Status determination
        if dc_row or status_raw in ['Vendido', 'Escriturado']:
            estado = 'Vendido'
        else:
            estado = 'Disponible'

        inmuebles.append({
            'id': inm_id,
            'proyecto_id': proyecto_id,
            'identificador': lot_name,
            'precio_venta': price,
            'canon_arriendo': 0,
            'estado': estado,
            'area_m2': 0,
            'matricula_inmobiliaria': None,
            'cedula_catastral': None,
            'linderos': f'Lote {lot_name} ubicado en el Paraje de Rio Negro, Vereda Paraderito, Proyecto Santa Isabel, Villavicencio - Meta.',
            'tradicion': 'Propiedad adquirida por Inversiones Ingeniería GC S.A.S.'
        })

    print(f"Total inmuebles prepared: {len(inmuebles)}")

    # 2. Process Clients & Contracts & Contratantes
    print("Processing DetalleCliente...")
    clients_map = {} # doc -> client dict
    contratos = []
    contratantes = []
    lot_to_contrato_id = {}

    for r in rows_dc[1:]:
        lot_name = str(r[hdr_dc.index('Modulo')]).strip()
        inm_id = lot_to_inmueble_id.get(lot_name)
        if not inm_id:
            print(f"Warning: Lot {lot_name} not found in Module!")
            continue

        contrato_id = make_uuid(f'contrato_santa_isabel_{lot_name}')
        lot_to_contrato_id[lot_name] = contrato_id

        # Buyer 1
        t1_name = clean_name(r[hdr_dc.index('Titular1')])
        t1_doc = clean_doc(r[hdr_dc.index('Cedula')])
        t1_email = str(r[hdr_dc.index('Email')]).strip() if r[hdr_dc.index('Email')] else ''
        t1_tel = clean_phone(r[hdr_dc.index('Telefono')])
        t1_ciudad = str(r[hdr_dc.index('Ciudad')]).strip() if r[hdr_dc.index('Ciudad')] else ''
        
        # Enrich from bitacora
        bit_info = bit_clients.get(t1_doc, {})
        t1_dir = bit_info.get('direccion') or (f"Ciudad: {t1_ciudad}" if t1_ciudad else 'Villavicencio')
        if not t1_email and bit_info.get('correo'):
            t1_email = bit_info['correo']
        if not t1_email:
            t1_email = f"cliente_{t1_doc}@inversionesgc.com"
        if not t1_tel and bit_info.get('celular'):
            t1_tel = bit_info['celular']

        if t1_doc not in clients_map:
            t1_id = make_uuid(f'cliente_{t1_doc}')
            clients_map[t1_doc] = {
                'id': t1_id,
                'tipo_persona': 'Natural',
                'documento': t1_doc,
                'nombre_razon_social': t1_name,
                'correo': t1_email,
                'telefono': t1_tel or '3115814360',
                'direccion': t1_dir
            }
        else:
            t1_id = clients_map[t1_doc]['id']

        # Add Contratante Principal
        contratantes.append({
            'id': make_uuid(f'contratante_{contrato_id}_{t1_doc}_principal'),
            'contrato_id': contrato_id,
            'cliente_id': t1_id,
            'rol_contratante': 'Comprador Principal'
        })

        # Co-buyer 2
        t2_name = clean_name(r[hdr_dc.index('Titular2')]) if r[hdr_dc.index('Titular2')] else ''
        t2_doc = clean_doc(r[hdr_dc.index('Cedula2')]) if r[hdr_dc.index('Cedula2')] else ''
        if t2_name and t2_doc:
            t2_email = str(r[hdr_dc.index('Email2')]).strip() if r[hdr_dc.index('Email2')] else f"cliente_{t2_doc}@inversionesgc.com"
            t2_tel = clean_phone(r[hdr_dc.index('Telefono2')]) or '3115814360'
            if t2_doc not in clients_map:
                t2_id = make_uuid(f'cliente_{t2_doc}')
                clients_map[t2_doc] = {
                    'id': t2_id,
                    'tipo_persona': 'Natural',
                    'documento': t2_doc,
                    'nombre_razon_social': t2_name,
                    'correo': t2_email,
                    'telefono': t2_tel,
                    'direccion': t1_dir
                }
            else:
                t2_id = clients_map[t2_doc]['id']

            contratantes.append({
                'id': make_uuid(f'contratante_{contrato_id}_{t2_doc}_copropietario'),
                'contrato_id': contrato_id,
                'cliente_id': t2_id,
                'rol_contratante': 'Co-propietario'
            })

        # Co-buyer 3
        t3_name = clean_name(r[hdr_dc.index('Titular3')]) if r[hdr_dc.index('Titular3')] else ''
        t3_doc = clean_doc(r[hdr_dc.index('Cedula3')]) if r[hdr_dc.index('Cedula3')] else ''
        if t3_name and t3_doc:
            t3_email = str(r[hdr_dc.index('Email3')]).strip() if r[hdr_dc.index('Email3')] else f"cliente_{t3_doc}@inversionesgc.com"
            t3_tel = clean_phone(r[hdr_dc.index('Telefono3')]) or '3115814360'
            if t3_doc not in clients_map:
                t3_id = make_uuid(f'cliente_{t3_doc}')
                clients_map[t3_doc] = {
                    'id': t3_id,
                    'tipo_persona': 'Natural',
                    'documento': t3_doc,
                    'nombre_razon_social': t3_name,
                    'correo': t3_email,
                    'telefono': t3_tel,
                    'direccion': t1_dir
                }
            else:
                t3_id = clients_map[t3_doc]['id']

            contratantes.append({
                'id': make_uuid(f'contratante_{contrato_id}_{t3_doc}_copropietario'),
                'contrato_id': contrato_id,
                'cliente_id': t3_id,
                'rol_contratante': 'Co-propietario'
            })

        # Contract financials
        val_tot = parse_amount(r[hdr_dc.index('ValorTotal')])
        cuota_inicial_prog = parse_amount(r[hdr_dc.index('TotalCuotaInicialProgramada')]) or parse_amount(r[hdr_dc.index('CuotaInicialProgramada')])
        fecha_venta = parse_date(r[hdr_dc.index('FechaVenta')])

        contratos.append({
            'id': contrato_id,
            'empresa_id': empresa_id,
            'inmueble_id': inm_id,
            'tipo_contrato': 'Compraventa',
            'fecha_inicio': fecha_venta,
            'fecha_fin': None,
            'valor_total': val_tot,
            'tiene_intereses': False,
            'tasa_interes_corriente': None,
            'tasa_interes_mora': 3.0,
            'estado_firma': 'Firmado',
            'monto_cuota_inicial': cuota_inicial_prog,
            'numero_cuotas_iniciales': 1,
            'monto_cuota_ordinaria': 0, # Will update from plan_pagos
            'numero_cuotas_ordinarias': 0, # Will update from plan_pagos
            'dia_pago_mensual': 15
        })

    # Ensure unique emails for all clients (prevent duplicate key violations)
    seen_emails = set()
    for doc, c in clients_map.items():
        email = c['correo'].lower().strip()
        if not email or email in seen_emails:
            if '@' in email:
                prefix, domain = email.split('@', 1)
                email = f"{prefix}+{doc}@{domain}"
            else:
                email = f"cliente_{doc}@inversionesgc.com"
        seen_emails.add(email)
        c['correo'] = email

    print(f"Total unique clients prepared: {len(clients_map)}")
    print(f"Total contracts prepared: {len(contratos)}")
    print(f"Total contratantes prepared: {len(contratantes)}")

    # 3. Process PlanDePagos and Pagos with Waterfall Allocation
    print("Processing PlanDePagos and Pagos...")
    sheet_pdp = wb_smart['PlanDePagos']
    rows_pdp = list(sheet_pdp.iter_rows(values_only=True))
    hdr_pdp = rows_pdp[0]
    pdp_by_lot = defaultdict(list)
    for r in rows_pdp[1:]:
        lot = str(r[hdr_pdp.index('Inmueble')]).strip()
        f = parse_date(r[hdr_pdp.index('Fecha')])
        m = parse_amount(r[hdr_pdp.index('Monto')])
        t = str(r[hdr_pdp.index('TipoPago')]).strip() if r[hdr_pdp.index('TipoPago')] else 'Cuota'
        tit = str(r[hdr_pdp.index('Titulo')]).strip() if r[hdr_pdp.index('Titulo')] else ''
        pdp_by_lot[lot].append({
            'fecha': f,
            'monto': m,
            'tipo_raw': t,
            'titulo': tit
        })

    # Sort each lot's cuotas by date
    for lot in pdp_by_lot:
        pdp_by_lot[lot].sort(key=lambda x: x['fecha'])

    sheet_pag = wb_smart['Pagos']
    rows_pag = list(sheet_pag.iter_rows(values_only=True))
    hdr_pag = rows_pag[0]
    pag_by_lot = defaultdict(list)
    for r in rows_pag[1:]:
        lot = str(r[hdr_pag.index('Inmueble')]).strip()
        f = parse_date(r[hdr_pag.index('Fecha')])
        m = parse_amount(r[hdr_pag.index('Monto')])
        rec = str(r[hdr_pag.index('NumeroRecibo')]).strip() if r[hdr_pag.index('NumeroRecibo')] else 'S/N'
        pag_by_lot[lot].append({
            'fecha': f,
            'monto': m,
            'recibo': rec
        })

    # Sort payments by date
    for lot in pag_by_lot:
        pag_by_lot[lot].sort(key=lambda x: x['fecha'])

    plan_pagos_records = []
    pagos_bitacora_records = []

    # Map of contrato_id -> stats
    contrato_stats = {}

    for lot_name, contrato_id in lot_to_contrato_id.items():
        raw_cuotas = pdp_by_lot.get(lot_name, [])
        raw_pagos = pag_by_lot.get(lot_name, [])

        # Build plan_pagos objects
        cuota_objs = []
        for seq, c in enumerate(raw_cuotas, 1):
            c_fecha = c['fecha']
            cuota_id = make_uuid(f'plan_pago_{contrato_id}_{seq}_{c_fecha}')
            
            tipo = 'ORDINARIA'
            if 'Separacion' in c['tipo_raw'] or 'Separacion' in c['titulo']:
                tipo = 'CUOTA_INICIAL'
            elif 'Crédito' in c['tipo_raw'] or 'Extra' in c['titulo']:
                tipo = 'EXTRAORDINARIA'

            cuota_objs.append({
                'id': cuota_id,
                'contrato_id': contrato_id,
                'numero_cuota': seq,
                'fecha_vencimiento': c['fecha'],
                'monto_cuota': c['monto'],
                'monto_pagado': 0.0,
                'monto_interes_mora': 0.0,
                'tipo_cuota': tipo,
                'estado': 'Pendiente'
            })

        # Waterfall allocation
        c_idx = 0
        for p_idx, p in enumerate(raw_pagos):
            rem_pago = p['monto']
            p_date = p['fecha']
            recibo_str = f"Recibo # {p['recibo']}" if p['recibo'] != 'S/N' else 'Comprobante de Pago'

            while rem_pago > 0.001 and c_idx < len(cuota_objs):
                c = cuota_objs[c_idx]
                needed = c['monto_cuota'] - c['monto_pagado']
                if needed <= 0.001:
                    c_idx += 1
                    continue

                alloc = min(rem_pago, needed)
                c['monto_pagado'] = round(c['monto_pagado'] + alloc, 2)
                rem_pago = round(rem_pago - alloc, 2)

                c_id = c['id']
                pagos_bitacora_records.append({
                    'id': make_uuid(f'pago_bitacora_{contrato_id}_{p_idx}_{c_id}_{alloc}'),
                    'plan_pagos_id': c['id'],
                    'fecha_pago': p_date,
                    'monto_pagado': alloc,
                    'metodo_pago': 'Transferencia',
                    'soporte_url': recibo_str
                })

                if c['monto_pagado'] >= c['monto_cuota'] - 0.01:
                    c_idx += 1

            # If there's an overpayment balance, assign it to the last cuota if exists
            if rem_pago > 0.001 and len(cuota_objs) > 0:
                last_c = cuota_objs[-1]
                last_c['monto_pagado'] = round(last_c['monto_pagado'] + rem_pago, 2)
                pagos_bitacora_records.append({
                    'id': make_uuid(f'pago_bitacora_{contrato_id}_{p_idx}_over_{rem_pago}'),
                    'plan_pagos_id': last_c['id'],
                    'fecha_pago': p_date,
                    'monto_pagado': rem_pago,
                    'metodo_pago': 'Transferencia',
                    'soporte_url': f"{recibo_str} (Abono a capital / saldo a favor)"
                })

        # Determine final status for each cuota
        ordinarias = [c for c in cuota_objs if c['tipo_cuota'] == 'ORDINARIA']
        avg_ordinaria = (sum(c['monto_cuota'] for c in ordinarias) / len(ordinarias)) if ordinarias else 0.0

        day_pago = 15
        if cuota_objs:
            try:
                day_pago = int(cuota_objs[0]['fecha_vencimiento'].split('-')[2])
            except:
                day_pago = 15

        contrato_stats[contrato_id] = {
            'num_ordinarias': len(ordinarias),
            'monto_ordinaria': round(avg_ordinaria, 2),
            'dia_pago': day_pago
        }

        for c in cuota_objs:
            if c['monto_pagado'] >= c['monto_cuota'] - 0.01:
                c['estado'] = 'Pagado'
            else:
                if c['fecha_vencimiento'] < today:
                    c['estado'] = 'Vencido'
                else:
                    c['estado'] = 'Pendiente'
            plan_pagos_records.append(c)

    # Update contract stats
    for ct in contratos:
        stats = contrato_stats.get(ct['id'])
        if stats:
            ct['numero_cuotas_ordinarias'] = stats['num_ordinarias']
            ct['monto_cuota_ordinaria'] = stats['monto_ordinaria']
            ct['dia_pago_mensual'] = stats['dia_pago']

    print(f"Total plan_pagos prepared: {len(plan_pagos_records)}")
    print(f"Total pagos_bitacora prepared: {len(pagos_bitacora_records)}")

    payload = {
        'inmuebles': inmuebles,
        'clientes': list(clients_map.values()),
        'contratos': contratos,
        'contratantes_contrato': contratantes,
        'plan_pagos': plan_pagos_records,
        'pagos_bitacora': pagos_bitacora_records
    }

    out_file = 'scratch/santa_isabel_payload.json'
    print(f"Writing payload to {out_file}...")
    with open(out_file, 'w', encoding='utf-8') as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)

    print("Data preparation complete successfully!")
    print(f"Summary:")
    print(f"  - Inmuebles: {len(inmuebles)}")
    print(f"  - Clientes: {len(clients_map)}")
    print(f"  - Contratos: {len(contratos)}")
    print(f"  - Contratantes: {len(contratantes)}")
    print(f"  - Plan Pagos: {len(plan_pagos_records)}")
    print(f"  - Pagos Bitacora: {len(pagos_bitacora_records)}")

if __name__ == '__main__':
    main()
